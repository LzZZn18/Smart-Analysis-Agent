# -*- coding: utf-8 -*-
"""提取模板 docx 结构（段落/标题/表格），用于对齐格式。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from pathlib import Path
from docx import Document

BASE = Path(r"E:\aiexcel\docs")
TMPL = BASE / "【考核材料模板】第21组考核材料-(20种重要统计学方法-梁梓轩-李炫標-刘师铭-林学炜)" \
             / "第21组考核材料-(20种重要统计学方法-梁梓轩-李炫標-刘师铭-林学炜)"

def dump(name):
    p = TMPL / name
    if not p.exists():
        print(f"!! 不存在: {p}")
        return
    doc = Document(str(p))
    print(f"\n===== {name} =====")
    print("-- 正文(前60个块) --")
    n = 0
    for para in doc.paragraphs[:60]:
        style = para.style.name if para.style else ""
        txt = para.text.strip()
        if txt:
            print(f"[{style}] {txt[:80]}")
        n += 1
    for ti, table in enumerate(doc.tables[:3]):
        print(f"-- 表格{ti+1} ({len(table.rows)}行x{len(table.columns)}列) --")
        for row in table.rows[:6]:
            print(" | ".join(c.text.strip()[:20] for c in row.cells))

dump("0-任务书.docx")
