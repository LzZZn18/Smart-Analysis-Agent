"""报告生成器：把一次（或多轮会话）分析结果渲染为 Word(docx) 或 PDF 文档。

支持 python-docx 生成 docx；reportlab 生成 PDF。
接收"记录列表"（每条 = 一轮问答），可还原多轮追问的完整对话：
每轮包含问题、结论、结果数据表、图表类型、分析计划、生成代码，保证可复现。
若传入单条 dict 会自动包装成单轮列表，向后兼容。
"""
from __future__ import annotations

import io
import json
from datetime import datetime
from typing import Any, Dict, List, Union

# 单条记录的字段：可以直接用 json.loads 还原的字段
_JSON_FIELDS = ("plan", "output", "charts")


def _as_records(records: Union[Dict[str, Any], List[Dict[str, Any]]]) -> List[Dict[str, Any]]:
    """统一入参为记录列表。传入单条 dict 时包装为单元素列表。"""
    if isinstance(records, dict):
        return [records]
    return records or []


def _decode_output(result: Dict[str, Any]) -> List[dict]:
    """把可能为 JSON 字符串的 output 还原为列表。"""
    raw = result.get("output", "[]")
    if isinstance(raw, list):
        return raw
    try:
        return json.loads(str(raw or "[]"))
    except (json.JSONDecodeError, ValueError):
        return []


def _render_chart_image(output: List[dict], chart_type: str = "bar") -> bytes:
    """根据结果数据 + 图表类型, 用 matplotlib 自绘一张 PNG 图表图片。

    Args:
        output: 结果数据(list[dict])
        chart_type: bar / line / pie

    Returns:
        PNG 图片字节流; 无法绘制时返回 b""。
    """
    if not output or not isinstance(output[0], dict):
        return b""
    try:
        import matplotlib
        matplotlib.use("Agg")   # 无界面后端
        import matplotlib.pyplot as plt
        # 中文字体：优先用系统中文字体, 缺则用英文兜底
        from matplotlib import font_manager as fm

        # 尝试配置中文字体
        plt.rcParams["font.sans-serif"] = [
            "Microsoft YaHei", "SimHei", "PingFang SC", "Arial Unicode MS", "DejaVu Sans",
        ]
        plt.rcParams["axes.unicode_minus"] = False

        # 提取类别列与数值列(与 charting 类似的启发)
        keys = list(output[0].keys())
        cat_key, val_key = None, None
        for k in keys:
            if not all(_is_number(row.get(k)) for row in output):
                cat_key = k
                break
        for k in keys:
            if k != cat_key and all(_is_number(row.get(k)) for row in output):
                val_key = k
                break
        if not cat_key or not val_key:
            return b""
        cats = [_safestr(row.get(cat_key)) for row in output]
        vals = [_asnum(row.get(val_key)) for row in output]

        fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=110)
        if chart_type == "pie":
            ax.pie(vals, labels=cats, autopct="%.1f%%", startangle=90)
            ax.set_title(f"{val_key} 占比分布")
        elif chart_type == "line":
            ax.plot(cats, vals, marker="o", color="#4f6ef7")
            ax.set_title(f"{cat_key} 与 {val_key} 的关系")
        else:  # bar 默认
            ax.bar(cats, vals, color="#4f6ef7")
            ax.set_title(f"{cat_key} 与 {val_key} 的关系")
        fig.tight_layout()

        buf = io.BytesIO()
        fig.savefig(buf, format="png")
        plt.close(fig)
        return buf.getvalue()
    except Exception:
        return b""


def _is_number(v) -> bool:
    """判断值是否为数值。"""
    if v is None or str(v).strip() == "":
        return False
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


def _asnum(v) -> float:
    """安全转数值。"""
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _safestr(v) -> str:
    """安全转字符串。"""
    return "" if v is None else str(v)


def _rows_for_docx(output: List[dict]) -> tuple[List[str], List[List[str]]]:
    """把输出数据转成 (表头, 行) 用于表格。"""
    if not output or not isinstance(output[0], dict):
        return [], []
    headers = list(output[0].keys())
    rows = [[str(r.get(h, "")) for h in headers] for r in output[:100]]
    return headers, rows


def _add_table_docx(doc, output: List[dict]) -> None:
    """在 docx 文档中添加一张结果数据表。"""
    from docx import Document  # noqa: F401 (类型提示)
    headers, rows = _rows_for_docx(output)
    if not headers:
        doc.add_paragraph("（本轮未返回数据表）")
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v


# ------ docx 渲染 ------

def render_docx(records: Union[Dict[str, Any], List[Dict[str, Any]]]) -> bytes:
    """渲染 Word 文档，返回 docx 字节流。支持多轮会话记录。"""
    from docx import Document
    from docx.shared import Pt

    recs = _as_records(records)
    doc = Document()

    total = len(recs)
    last_q = recs[-1].get("question", "") if total else ""
    doc.add_heading(f"数据分析报告 — {last_q}", level=1)
    meta = doc.add_paragraph()
    meta.add_run(f"生成时间：{datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')}\n")
    meta.add_run(f"会话轮次：{total} 轮 | 报表编号：{recs[0].get('report_id', '') if recs else ''}")

    # 逐轮渲染
    for idx, rec in enumerate(recs, start=1):
        doc.add_heading(f"第 {idx} 轮 ｜ {rec.get('question', '')}", level=2)
        doc.add_paragraph(f"结论：{rec.get('conclusion', '（无）')}")
        _add_table_docx(doc, _decode_output(rec))

        # 图表类型 + 真实图表图片
        try:
            charts = json.loads(rec.get("charts", "[]") or "[]")
        except json.JSONDecodeError:
            charts = []
        if charts:
            types = "、".join(c.get("chart_type", "未知") for c in charts)
            doc.add_paragraph(f"图表类型：{types}（可按生成代码溯源）")
            output = _decode_output(rec)
            if output:
                img_bytes = _render_chart_image(output, charts[0].get("chart_type", "bar"))
                if img_bytes:
                    from docx.shared import Inches
                    doc.add_picture(io.BytesIO(img_bytes), width=Inches(5))
        else:
            doc.add_paragraph("（本轮未生成图表）")

        # 该轮计划与代码（可复现）
        try:
            plan = json.loads(rec.get("plan", "[]") or "[]")
        except json.JSONDecodeError:
            plan = []
        doc.add_paragraph("分析计划：" + " → ".join(str(p) for p in plan))
        code = doc.add_paragraph()
        run = code.add_run(rec.get("code", ""))
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------ PDF 渲染 ------

def render_pdf(records: Union[Dict[str, Any], List[Dict[str, Any]]]) -> bytes:
    """渲染 PDF 文档，返回字节流（reportlab 实现，支持中文）。支持多轮会话。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    recs = _as_records(records)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="数据分析报告")
    styles = getSampleStyleSheet()

    # 注册中文字体（内置 CID 字体, 无需外部文件）
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        zh = "STSong-Light"
    except Exception:
        zh = "Helvetica"

    body_style = ParagraphStyle("zh", parent=styles["BodyText"], fontName=zh, fontSize=10, leading=15)
    h1 = ParagraphStyle("zh1", parent=styles["Heading1"], fontName=zh, fontSize=16)
    h2 = ParagraphStyle("zh2", parent=styles["Heading2"], fontName=zh, fontSize=13)

    story: list = []
    total = len(recs)
    last_q = recs[-1].get("question", "") if total else ""
    story.append(Paragraph(f"数据分析报告 — {last_q}", h1))
    story.append(Spacer(1, 6))
    story.append(Paragraph(f"生成时间：{datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')}", body_style))
    story.append(Paragraph(f"会话轮次：{total} 轮", body_style))
    story.append(Spacer(1, 10))

    for idx, rec in enumerate(recs, start=1):
        story.append(Paragraph(f"第 {idx} 轮 ｜ {rec.get('question', '')}", h2))
        story.append(Paragraph(f"结论：{rec.get('conclusion', '（无）')}", body_style))

        # 数据表
        output = _decode_output(rec)
        headers, rows = _rows_for_docx(output)
        if headers:
            t = Table([headers] + rows, repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e0e0e0")),
                ("FONTNAME", (0, 0), (-1, -1), zh),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]))
            story.append(t)
        else:
            story.append(Paragraph("（本轮未返回数据表）", body_style))

        # 图表类型 + 真实图表图片
        try:
            charts = json.loads(rec.get("charts", "[]") or "[]")
        except json.JSONDecodeError:
            charts = []
        if charts:
            types = "、".join(c.get("chart_type", "未知") for c in charts)
            story.append(Paragraph(f"图表类型：{types}（可按生成代码溯源）", body_style))
            # 嵌入真实图表图片(基于结果数据自绘)
            if output:
                img_bytes = _render_chart_image(output, charts[0].get("chart_type", "bar"))
                if img_bytes:
                    story.append(Image(io.BytesIO(img_bytes), width=400, height=225))
        else:
            story.append(Paragraph("（本轮未生成图表）", body_style))

        # 计划与代码
        try:
            plan = json.loads(rec.get("plan", "[]") or "[]")
        except json.JSONDecodeError:
            plan = []
        story.append(Paragraph("分析计划：" + " → ".join(str(p) for p in plan), body_style))
        story.append(Spacer(1, 3))
        story.append(Paragraph("<pre>" + rec.get("code", "") + "</pre>", body_style))
        story.append(Spacer(1, 8))

    doc.build(story)
    return buf.getvalue()
