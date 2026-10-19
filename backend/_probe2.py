# -*- coding: utf-8 -*-
"""批量提取多个模板 docx 结构，用于对齐格式。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from pathlib import Path
from docx import Document

TMPL = Path(r"E:\aiexcel\docs") / "【考核材料模板】第21组考核材料-(20种重要统计学方法-梁梓轩-李炫標-刘师铭-林学炜)" \
        / "第21组考核材料-(20种重要统计学方法-梁梓轩-李炫標-刘师铭-林学炜)"

def dump(name, maxpara=80):
    p = TMPL / name
    if not p.exists():
        print(f"!! 不存在: {p}"); return
    if p.suffix.lower() == ".doc":
        print(f"\n===== {name} =====\n(.doc 旧格式, python-docx 无法读取)")
        return
    doc = Document(str(p))
    print(f"\n===== {name} =====")
    print("-- 段落(前80块) --")
    for para in doc.paragraphs[:maxpara]:
        style = para.style.name if para.style else ""
        txt = para.text.strip()
        if txt:
            # 判断是否是标题样式
            flag = "H" if ("Heading" in style or style.startswith("标题")) else " " * 2 + style[:8]
            print(f"[{flag}] {txt[:70]}")
    for ti, table in enumerate(doc.tables):
        print(f"-- 表格{ti+1} ({len(table.rows)}r) --")
        for row in table.rows[:8]:
            print("   ", " | ".join(c.text.strip()[:18] for c in row.cells))

dump("0-任务书.docx")
dump("1-项目提案.docx")
