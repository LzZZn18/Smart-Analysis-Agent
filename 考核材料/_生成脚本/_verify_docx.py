# -*- coding: utf-8 -*-
"""校验生成的 docx 可打开、结构正确。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from pathlib import Path
from docx import Document

D = Path(__file__).parent / "docx"
files = sorted(D.glob("*.docx"))
print(f"共 {len(files)} 个 docx")
bad = 0
for p in files:
    try:
        d = Document(str(p))
        paragraphs = [x.text.strip() for x in d.paragraphs if x.text.strip()]
        ntable = len(d.tables)
        first = paragraphs[0] if paragraphs else "(空)"
        print(f"  OK {p.name}: 段落{len(paragraphs)} 表格{ntable} | 首行: {first[:30]}")
    except Exception as exc:
        bad += 1
        print(f"  FAIL {p.name}: {exc}")
print("=== 校验" + ("全部通过" if bad == 0 else f"{bad} 个失败") + " ===")
